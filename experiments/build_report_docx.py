"""把實驗報告 Markdown 原稿轉成可編輯的 Word 文件。"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "OnboardBot_中文實驗報告.md"
OUTPUT = ROOT / "OnboardBot_中文實驗報告.docx"
FONT = "Microsoft JhengHei"


def set_font(run, size=None, bold=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        tag = "w:" + side
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", 90), ("left", 90), ("bottom", 90), ("right", 90)):
        item = OxmlElement("w:" + side)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    tc_pr.append(margins)


def inline_text(paragraph, text, size=None):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))")
    for token in pattern.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            set_font(paragraph.add_run(token[2:-2]), size, True)
        elif token.startswith("`") and token.endswith("`"):
            set_font(paragraph.add_run(token[1:-1]), size, name="Consolas")
        elif token.startswith("[") and "](" in token:
            label, url = token[1:-1].split("](", 1)
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "000000")
            properties.append(color)
            run.append(properties)
            text_element = OxmlElement("w:t")
            text_element.text = label
            run.append(text_element)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        else:
            set_font(paragraph.add_run(token), size)


def add_table(doc, lines):
    rows = [[item.strip() for item in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell or "") for cell in row)]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    widths = [Inches(value) for value in (0.48, 1.62, 0.72, 2.88, 1.05)]
    for ri, data in enumerate(rows):
        for ci, content in enumerate(data):
            cell = table.cell(ri, ci)
            cell.width = widths[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if ri == 0:
                set_cell_shading(cell, "DCE6F1")
            elif ri % 2 == 0:
                set_cell_shading(cell, "F7F9FB")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            if ci in (0, 2, 4):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline_text(p, content, 8.7)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
    for cell in table.rows[0].cells:
        cell._tc.get_or_add_tcPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.87)
    section.right_margin = Inches(0.87)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28
    for name, size, before, after in (("Title", 17.5, 0, 12), ("Heading 1", 13, 14, 7), ("Heading 2", 11.2, 11, 5)):
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.underline = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        ppr = style.element.get_or_add_pPr()
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("OnboardBot 實驗報告"), 8.5)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("第 "), 8.5)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    set_font(footer.add_run(" 頁"), 8.5)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            inline_text(p, line[2:])
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            inline_text(p, line[3:])
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            inline_text(p, line[4:])
        elif line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        elif line.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.1
                set_font(p.add_run(lines[i]), 8.6, name="Consolas")
                i += 1
        else:
            p = doc.add_paragraph()
            if line.startswith("**實驗") or line.startswith("**報告"):
                p.paragraph_format.space_after = Pt(2)
            inline_text(p, line.rstrip(" "))
        i += 1

    doc.core_properties.title = "OnboardBot 中文實驗報告"
    doc.core_properties.subject = "LangGraph 與 Tavily MCP 系統實測"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
